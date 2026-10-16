"""The World Bank Pink Sheet reader: many commodities, many units, one file."""

import io

import openpyxl
import pytest

from merx.sources import worldbank


def pink_sheet(
    names=("Copper", "Gold"),
    units=("($/mt)", "($/troy oz)"),
    rows=(("1960M01", 607.0, 35.0), ("1960M02", 610.5, None)),
    release="Updated on October 06, 2026",
    sheet=worldbank.SHEET,
) -> bytes:
    """A workbook laid out like the real Pink Sheet (verified 2026-10-16)."""
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = sheet
    ws.cell(row=4, column=1, value=release)
    for i, name in enumerate(names, start=2):
        ws.cell(row=worldbank.NAME_ROW, column=i, value=name)
    for i, unit in enumerate(units, start=2):
        ws.cell(row=worldbank.UNIT_ROW, column=i, value=unit)
    for r, row in enumerate(rows, start=worldbank.FIRST_DATA_ROW):
        ws.cell(row=r, column=1, value=row[0])
        for i, value in enumerate(row[1:], start=2):
            # The real file writes a placeholder glyph, not a blank, where a price is
            # not available.
            ws.cell(row=r, column=i, value="…" if value is None else value)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_a_named_column_becomes_a_monthly_series():
    parsed = worldbank.parse(pink_sheet(), {"price": "Copper"})
    assert [(d.isoformat(), v) for _, d, v in parsed.observations] == [
        ("1960-01-01", 607.0),
        ("1960-02-01", 610.5),
    ]


def test_the_placeholder_glyph_is_a_gap_not_a_zero():
    """Reading the glyph as 0 would put a free tonne of gold in the history."""
    parsed = worldbank.parse(pink_sheet(), {"price": "Gold"})
    assert [v for _, _, v in parsed.observations] == [35.0, None]


def test_the_file_states_its_own_units():
    parsed = worldbank.parse(pink_sheet(), {"copper": "Copper", "gold": "Gold"})
    assert parsed.units == {"copper": "USD/mt", "gold": "USD/troy oz"}


def test_the_release_date_is_read():
    parsed = worldbank.parse(pink_sheet(), {"price": "Copper"})
    assert parsed.release_date.isoformat() == "2026-10-06"


def test_a_column_that_is_gone_fails_loudly():
    with pytest.raises(worldbank.PinkSheetError, match="Lithium"):
        worldbank.parse(pink_sheet(), {"price": "Lithium"})


def test_a_renamed_sheet_fails_loudly():
    with pytest.raises(worldbank.PinkSheetError, match="sheet"):
        worldbank.parse(pink_sheet(sheet="Prices"), {"price": "Copper"})


def test_a_repriced_benchmark_is_visible_in_the_units():
    """$/mt to $/kg would rescale the series by a thousand and still look plausible.
    The contract's declared unit is compared against this, in the reader."""
    parsed = worldbank.parse(pink_sheet(units=("($/kg)", "($/troy oz)")), {"price": "Copper"})
    assert parsed.units == {"price": "USD/kg"}
